"""
Genera el Excel de revision de leads en "Salidas express/".
Fuente: data/leads.csv + data/apify-ledger.jsonl (ambos gitignored).

    python Scripts/gen_excel_leads.py
"""

import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

REPO = Path(__file__).resolve().parents[1]
DATA = REPO / "data"
SALIDA = REPO / "Salidas express"

# Paleta Proaudita (CONTEXT.md)
AZUL = "0019FF"
NAVY = "1A1A5E"
GRIS = "F5F7FA"
TEXTO = "0D0D2B"

H_FILL = PatternFill("solid", fgColor=NAVY)
H_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BORDE = Border(*[Side(style="thin", color="D9D9D9")] * 4)


def leer_leads() -> list:
    with (DATA / "leads.csv").open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def leer_ledger() -> list:
    p = DATA / "apify-ledger.jsonl"
    if not p.exists():
        return []
    return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]


def estilar_encabezado(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        celda = ws.cell(row=1, column=c)
        celda.fill = H_FILL
        celda.font = H_FONT
        celda.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def autoancho(ws, maximo: int = 46) -> None:
    for col in ws.columns:
        largo = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(largo + 2, maximo)


def hoja_leads(wb: Workbook, leads: list) -> None:
    ws = wb.create_sheet("Leads")
    cols = [
        ("tier", "Tier"), ("prioridad", "Prioridad"), ("nombre", "Empresa"),
        ("familia", "Familia"), ("rubro", "Rubro"), ("ciudad", "Ciudad"),
        ("email", "Email"), ("email_nominal", "Llega a persona"),
        ("telefono", "Telefono"), ("sitio_web", "Sitio web"),
        ("rating", "Rating"), ("n_reviews", "Resenas"), ("tamano_proxy", "Tamano (proxy)"),
        ("score", "Score dato"), ("senal", "Senal para el primer contacto"),
        ("estado", "Estado"), ("fuente", "Fuente"),
        ("contacto_nombre", "Contacto"), ("contacto_cargo", "Cargo"),
        ("fecha_extraccion", "Extraido"),
        ("fecha_email_1", "Email 1"), ("fecha_email_2", "Email 2"), ("notas", "Notas"),
    ]
    ws.append([t for _, t in cols])

    for l in sorted(leads, key=lambda x: (-int(x.get("prioridad") or 0), -int(x.get("score") or 0))):
        fila = []
        for k, _ in cols:
            v = l.get(k, "")
            if k in ("score", "n_reviews", "prioridad"):
                v = int(v or 0)
            elif k == "rating":
                v = float(v) if v else None
            fila.append(v)
        ws.append(fila)

    n = ws.max_row
    estilar_encabezado(ws, len(cols))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"

    # Semaforo por tier: A verde, B ambar, C gris
    for letra, relleno, color in (("A", "C6EFCE", "006100"),
                                  ("B", "FFEB9C", "9C5700"),
                                  ("C", "EDEDED", "6B7280")):
        ws.conditional_formatting.add(f"A2:B{n}", CellIsRule(
            operator="equal", formula=[f'"{letra}"'],
            fill=PatternFill("solid", fgColor=relleno), font=Font(color=color, bold=True)))

    idx = {k: i for i, (k, _) in enumerate(cols)}
    for fila in ws.iter_rows(min_row=2, max_row=n, max_col=len(cols)):
        for c in fila:
            c.border = BORDE
            c.alignment = Alignment(vertical="center")
        for k in ("tier", "prioridad", "score", "n_reviews"):
            fila[idx[k]].alignment = Alignment(horizontal="center", vertical="center")
        # Email y sitio como hipervinculos: la revision se hace clickeando
        if fila[idx["email"]].value:
            fila[idx["email"]].hyperlink = f"mailto:{fila[idx['email']].value}"
            fila[idx["email"]].font = Font(color=AZUL, underline="single")
        if fila[idx["sitio_web"]].value:
            fila[idx["sitio_web"]].hyperlink = str(fila[idx["sitio_web"]].value)
            fila[idx["sitio_web"]].font = Font(color=AZUL, underline="single")
    autoancho(ws)


def hoja_resumen(wb: Workbook, leads: list, ledger: list) -> None:
    ws = wb.create_sheet("Resumen", 0)
    gasto = sum(d["usd"] for d in ledger)
    lugares = sum(d["lugares"] for d in ledger)
    n = len(leads)

    ws["A1"] = "Proaudita — Prospeccion outbound"
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = f"Scraping Google Maps via Apify · generado {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(size=10, italic=True, color="6B7280")

    def bloque(fila: int, titulo: str, pares: list) -> int:
        ws.cell(row=fila, column=1, value=titulo).font = Font(bold=True, size=12, color=AZUL)
        fila += 1
        for k, v in pares:
            ws.cell(row=fila, column=1, value=k).font = Font(bold=False)
            c = ws.cell(row=fila, column=2, value=v)
            c.font = Font(bold=True, color=TEXTO)
            c.alignment = Alignment(horizontal="left")
            fila += 1
        return fila + 1

    f = bloque(4, "Presupuesto", [
        ("Autorizado (USD)", 5.00),
        ("Gastado (USD)", round(gasto, 4)),
        ("Saldo (USD)", round(5.00 - gasto, 4)),
        ("Runs ejecutados", len(ledger)),
        ("Runs exitosos", sum(1 for d in ledger if not d.get("nota", "").startswith("status="))),
    ])

    f = bloque(f, "Rendimiento", [
        ("Lugares scrapeados", lugares),
        ("Leads de Apify con correo", sum(1 for l in leads if l.get("fuente") == "apify-google-maps")),
        ("Tasa lugar -> lead", f"{sum(1 for l in leads if l.get('fuente') == 'apify-google-maps') / lugares:.1%}" if lugares else "-"),
        ("Costo por lead de Apify (USD)", round(gasto / max(sum(1 for l in leads if l.get("fuente") == "apify-google-maps"), 1), 4)),
        ("Leads de gremios (sin costo)", sum(1 for l in leads if l.get("fuente", "").startswith(("cchc", "alog", "clinicas")))),
        ("Costo por lugar (USD)", 0.006),
    ])

    f = bloque(f, "Cola de envio (a 20 correos/dia habil)", [
        ("Tier A — escribir primero", sum(1 for l in leads if l.get("tier") == "A")),
        ("Tier B", sum(1 for l in leads if l.get("tier") == "B")),
        ("Tier C — solo si sobra capacidad", sum(1 for l in leads if l.get("tier") == "C")),
        ("Sin correo — por enriquecer", sum(1 for l in leads if l.get("estado") == "por_enriquecer")),
        ("Dias habiles para agotar la cola", round(sum(1 for l in leads if l["email"]) / 20, 1)),
    ])

    f = bloque(f, "Calidad de la muestra", [
        ("Email que llega a una persona", sum(1 for l in leads if l.get("email_nominal") == "si")),
        ("Email en dominio propio .cl", sum(1 for l in leads if l["email"].lower().endswith(".cl"))),
        ("Con senal para primer contacto", sum(1 for l in leads if l["senal"] != "needs-research")),
        ("Sin senal (needs-research)", sum(1 for l in leads if l["senal"] == "needs-research")),
    ])

    ws.cell(row=f, column=1, value="Ojo al revisar").font = Font(bold=True, size=12, color=AZUL)
    for i, nota in enumerate([
        "El score es un proxy de calidad, no una verificacion de tamano. Google Maps no",
        "publica facturacion: un score 95 puede ser una empresa mas chica de lo esperado.",
        "Se filtraron cadenas multinacionales, correos de reservas y micro-empresas, pero",
        "la validacion final de si califica como cliente ideal ($500M-$5B CLP) es humana.",
        "Marcar en la columna Notas las que NO sirven para reentrenar los filtros.",
    ], start=1):
        ws.cell(row=f + i, column=1, value=nota).font = Font(size=10, color="6B7280")

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18


def hoja_segmentos(wb: Workbook, leads: list) -> None:
    ws = wb.create_sheet("Segmentos")
    fila = 1
    for titulo, clave in [("Por familia de rubro", "familia"), ("Por fuente", "fuente"), ("Por ciudad", "ciudad")]:
        ws.cell(row=fila, column=1, value=titulo).font = Font(bold=True, size=12, color=AZUL)
        fila += 1
        ws.cell(row=fila, column=1, value="Segmento").font = H_FONT
        ws.cell(row=fila, column=1).fill = H_FILL
        ws.cell(row=fila, column=2, value="Leads").font = H_FONT
        ws.cell(row=fila, column=2).fill = H_FILL
        ws.cell(row=fila, column=3, value="Prioridad prom.").font = H_FONT
        ws.cell(row=fila, column=3).fill = H_FILL
        fila += 1
        cnt = Counter((l[clave] or "sin dato").split(",")[0].strip() for l in leads)
        for seg, k in cnt.most_common():
            proms = [int(l["prioridad"] or 0) for l in leads
                     if (l[clave] or "sin dato").split(",")[0].strip() == seg]
            ws.cell(row=fila, column=1, value=seg)
            ws.cell(row=fila, column=2, value=k)
            ws.cell(row=fila, column=3, value=round(sum(proms) / len(proms), 1))
            fila += 1
        fila += 2
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 12


def hoja_gasto(wb: Workbook, ledger: list) -> None:
    ws = wb.create_sheet("Gasto Apify")
    ws.append(["Fecha", "Run ID", "Tanda", "USD", "Lugares", "Leads brutos", "USD/lead"])
    for d in ledger:
        lc = d["leads_calificados"]
        ws.append([d["ts"][:16].replace("T", " "), d["run_id"], d.get("nota", ""),
                   round(d["usd"], 4), d["lugares"], lc,
                   round(d["usd"] / lc, 4) if lc else None])
    n = ws.max_row
    ws.append([])
    ws.append(["TOTAL", "", "", round(sum(d["usd"] for d in ledger), 4),
               sum(d["lugares"] for d in ledger), sum(d["leads_calificados"] for d in ledger), ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color=NAVY)
    estilar_encabezado(ws, 7)
    for fila in ws.iter_rows(min_row=2, max_row=n, max_col=7):
        for c in fila:
            c.border = BORDE
    ws["A" + str(ws.max_row + 2)] = ("Leads brutos = antes de la purga manual de cadenas "
                                     "multinacionales y micro-empresas (13 descartados).")
    ws["A" + str(ws.max_row)].font = Font(size=10, italic=True, color="6B7280")
    autoancho(ws)


def main() -> None:
    leads, ledger = leer_leads(), leer_ledger()
    wb = Workbook()
    wb.remove(wb.active)
    hoja_leads(wb, leads)
    hoja_segmentos(wb, leads)
    hoja_gasto(wb, ledger)
    hoja_resumen(wb, leads, ledger)   # se inserta como primera hoja
    wb.active = 0

    SALIDA.mkdir(parents=True, exist_ok=True)
    destino = SALIDA / f"leads-proaudita-{datetime.today():%Y-%m-%d}.xlsx"
    try:
        wb.save(destino)
    except PermissionError:
        # El archivo esta abierto en Excel: Windows bloquea la escritura.
        destino = destino.with_name(f"{destino.stem}-{datetime.now():%H%M}.xlsx")
        wb.save(destino)
        print("AVISO: el archivo anterior estaba abierto en Excel; se guardo una version nueva.")
    print(f"OK -> {destino}  ({len(leads)} leads, {len(ledger)} runs)")


if __name__ == "__main__":
    main()
